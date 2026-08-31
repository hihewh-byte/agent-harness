"""Helpers for broker xlsx sheets with preamble rows before CSV headers."""

from __future__ import annotations

import csv
import io
from typing import Any

from tax_agent.parser.fidelity_parser import find_fidelity_header_row
from tax_agent.parser.futu_parser import find_futu_header_row
from tax_agent.parser.schwab_parser import find_schwab_header_row

# 富途税表 xlsx 在 read_only 模式下常只读出首列，须用密集读取
FUTU_TAX_SHEET_MARKERS = frozenset(
    {"证券-交易流水", "股息、利息及其他收入", "账户信息", "参考汇率"}
)


def row_to_csv_line(row: tuple[Any, ...]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["" if c is None else c for c in row])
    return buf.getvalue().rstrip("\r\n")


def rows_to_text_lines(rows: list[tuple[Any, ...]], limit: int = 80) -> list[str]:
    return [row_to_csv_line(r) for r in rows[:limit]]


def find_header_in_sheet_rows(
    rows: list[tuple[Any, ...]],
) -> tuple[int, list[str]] | None:
    if not rows:
        return None
    lines = rows_to_text_lines(rows)
    return (
        find_fidelity_header_row(lines)
        or find_futu_header_row(lines)
        or find_schwab_header_row(lines)
    )


def load_workbook_sheet_data(data: bytes) -> tuple[list[str], dict[str, list[tuple[Any, ...]]]]:
    """Load all sheets with dense cell reads (fixes merged/sparse read_only rows)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheet_names = list(wb.sheetnames)
    rows_by_sheet: dict[str, list[tuple[Any, ...]]] = {}
    for name in sheet_names:
        ws = wb[name]
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        if max_row == 0 or max_col == 0:
            rows_by_sheet[name] = []
            continue
        sheet_rows: list[tuple[Any, ...]] = []
        for r in range(1, max_row + 1):
            sheet_rows.append(tuple(ws.cell(r, c).value for c in range(1, max_col + 1)))
        rows_by_sheet[name] = sheet_rows
    wb.close()
    return sheet_names, rows_by_sheet


def is_futu_tax_workbook(sheet_names: list[str]) -> bool:
    """Any Futu tax-package sheet marker, or a sheet named like 证券-交易流水."""
    names = set(sheet_names)
    if any("交易流水" in n for n in names):
        return True
    if "股息、利息及其他收入" in names and "账户信息" in names:
        return True
    return bool(names & FUTU_TAX_SHEET_MARKERS)


def sheet_rows_to_csv_text(
    rows: list[tuple[Any, ...]],
    header_idx: int,
    header_cells: list[str],
) -> str:
    """Rebuild CSV text from header row onward for broker CSV parsers."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(header_cells)
    for row in rows[header_idx + 1 :]:
        if row is None:
            continue
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        writer.writerow(["" if c is None else c for c in row])
    return out.getvalue()

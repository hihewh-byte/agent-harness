#!/usr/bin/env python3
"""Build minimal IBKR-style xlsx for parser tests."""

from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("pip install openpyxl")

out = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "broker" / "ibkr_mini.xlsx"
wb = openpyxl.Workbook()
wb.remove(wb.active)

div = wb.create_sheet("Dividends")
div.append(["Date", "Symbol", "Description", "Amount", "Tax", "Currency"])
div.append(["2024-06-15", "AAPL", "Cash Dividend", 1000, 300, "USD"])

tr = wb.create_sheet("Trades")
tr.append(["Date/Time", "Symbol", "Quantity", "Proceeds", "Realized P/L", "Comm/Fee", "Currency"])
tr.append(["2024-08-01", "AAPL", -10, 5000, 500, -1.25, "USD"])

wh = wb.create_sheet("Withholding Tax")
wh.append(["Date", "Description", "Amount", "Currency"])
wh.append(["2024-06-16", "NRA Tax", 50, "USD"])

out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print("wrote", out)

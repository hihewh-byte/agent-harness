#!/usr/bin/env python3
"""Anonymize Futu 2022 tax xlsx → tests/fixtures/broker/ (local dev only).

Sources (first match wins):
  1. FUTU_FIXTURE_SRC_ANNUAL / FUTU_FIXTURE_SRC_INCOME env vars
  2. docs/*_Annual_Statement_*.xlsx (gitignored — never committed)

Public clones use committed fixtures only; this script is optional for maintainers.
"""

from __future__ import annotations

import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT_ANNUAL = ROOT / "tests" / "fixtures" / "broker" / "futu_tax_2022_annual.xlsx"
OUT_INCOME = ROOT / "tests" / "fixtures" / "broker" / "futu_tax_2022_income_summary.xlsx"

REDACT_HEADERS = frozenset(
    {"姓名", "账户名称", "账户号码", "牛牛号", "Name", "Account Name", "Account Number"}
)


def _resolve_src_annual() -> Path | None:
    env = (os.environ.get("FUTU_FIXTURE_SRC_ANNUAL") or "").strip()
    if env:
        p = Path(env)
        return p if p.is_file() else None
    hits = sorted((ROOT / "docs").glob("*_Annual_Statement_*.xlsx"))
    return hits[0] if hits else None


def _resolve_src_income() -> Path | None:
    env = (os.environ.get("FUTU_FIXTURE_SRC_INCOME") or "").strip()
    if env:
        p = Path(env)
        return p if p.is_file() else None
    hits = sorted((ROOT / "docs").glob("*Interest*Income*.xlsx"))
    return hits[0] if hits else None


def _redact(val) -> object:
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return val
    if re.fullmatch(r"\d{8,}", s):
        return "1000000000000001"
    if re.search(r"帳戶|账户|Account", s):
        return "美股融资账户(0001)"
    if len(s) <= 4 and re.search(r"[\u4e00-\u9fff]", s):
        return "测试用户"
    return val


def anonymize_workbook(src: Path, dest: Path) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(src)
    changed = 0
    for ws in wb.worksheets:
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        headers = [ws.cell(1, c).value for c in range(1, max_col + 1)]
        redact_cols = {
            i + 1
            for i, h in enumerate(headers)
            if h and str(h).strip() in REDACT_HEADERS
        }
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                cell = ws.cell(r, c)
                if cell.value is None:
                    continue
                if r == 1 and str(cell.value).strip() in REDACT_HEADERS:
                    continue
                if c in redact_cols or (r > 1 and c in redact_cols):
                    new = _redact(cell.value)
                    if new != cell.value:
                        cell.value = new
                        changed += 1
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    wb.close()
    return changed


def main() -> int:
    src_annual = _resolve_src_annual()
    if not src_annual:
        print(
            "missing source annual xlsx — set FUTU_FIXTURE_SRC_ANNUAL or place file under docs/",
            file=sys.stderr,
        )
        return 1
    n1 = anonymize_workbook(src_annual, OUT_ANNUAL)
    print(f"wrote {OUT_ANNUAL} ({n1} cells redacted)")
    src_income = _resolve_src_income()
    if src_income:
        n2 = anonymize_workbook(src_income, OUT_INCOME)
        print(f"wrote {OUT_INCOME} ({n2} cells redacted)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Minimal Futu income-summary xlsx with USD dividend/interest for SSOT selfchecks."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tests" / "fixtures" / "broker" / "futu_classified_income_usd.xlsx"


def main() -> int:
    import openpyxl

    wb = openpyxl.Workbook()
    ws_info = wb.active
    ws_info.title = "账户信息"
    ws_info.append(["姓名", "牛牛号", "账户名称", "年份"])
    ws_info.append(["测试用户", "1000000000000001", "美股融资账户(0001)", "2022"])

    ws_inc = wb.create_sheet("股息、利息及其他收入")
    ws_inc.append(["牛牛号", "年份", "账户名称", "全年股息", "全年股息税", "全年利息", "币种"])
    ws_inc.append(["1000000000000001", "2022", "美股融资账户(0001)", "120.50", "18.08", "45.25", "USD"])

    ws_fx = wb.create_sheet("参考汇率")
    ws_fx.append(["note", "USD"])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    wb.close()
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

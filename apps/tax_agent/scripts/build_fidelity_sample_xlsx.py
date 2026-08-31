#!/usr/bin/env python3
"""Build Fidelity-style multi-sheet xlsx (History + Gains/Losses with preamble)."""

from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "tests" / "fixtures" / "broker"
HIST = FIX / "fidelity_activity_realistic.csv"
GL = FIX / "fidelity_gains_losses.csv"
OUT = FIX / "fidelity_multi_sheet.xlsx"


def _load_lines(path: Path) -> list[list[str]]:
    import csv

    lines: list[list[str]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        lines.append(next(csv.reader([line])))
    return lines


def _write_sheet(ws, rows: list[list[str]]) -> None:
    for row in rows:
        ws.append(row)


def main() -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb.create_sheet("Account History"), _load_lines(HIST))
    _write_sheet(wb.create_sheet("Gains and Losses"), _load_lines(GL))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
